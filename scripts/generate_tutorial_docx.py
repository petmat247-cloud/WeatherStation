#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generátor Word dokumentu (.docx): Návod na obnovu Weathercloud Cookie
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Nastaví barvu pozadí buňky v tabulce."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Nastaví vnitřní odsazení buňky."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    """Vloží blok kódu s šedým pozadím a písmem Consolas."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Rámeček
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    # Prázdný odstavec po bloku
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_callout(doc, text, title="UPOZORNĚNÍ / TIP", callout_type="info"):
    """Vloží zvýrazněný infobox (tip, info, warning)."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    border_color = "2563EB" if callout_type == "info" else ("16A34A" if callout_type == "success" else "EA580C")
    bg_color = "EFF6FF" if callout_type == "info" else ("F0FDF4" if callout_type == "success" else "FFF7ED")
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="32" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor.from_string(border_color)
    
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def build_tutorial_doc():
    doc = Document()
    
    # Nastavení okrajů (normální 2.5 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    # ── TITULNÍ BLOK ──────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Návod: Obnova Weathercloud Cookie")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(24)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Tmavě modrá
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Krok za krokem: Jak obnovit načítání vnitřní teploty a vlhkosti z domova pro meteorologický dashboard Stanice")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(11.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    # Oddělovací linka
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(16)
    p_line_border = parse_xml(f'''
        <w:pBdr {nsdecls("w")}>
            <w:bottom w:val="single" w:sz="12" w:space="1" w:color="CBD5E1"/>
        </w:pBdr>
    ''')
    p_line._element.get_or_add_pPr().append(p_line_border)
    
    # ── 1. PROČ JE TO POTŘEBA ────────────────────────────────────────────────
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run("1. Proč a kdy je potřeba cookie obnovit?")
    r.font.name = "Segoe UI"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Meteostanice Sencor SWS 9898 WIFI měří jak venkovní podmínky, tak vnitřní teplotu a vlhkost v domě. "
        "Zatímco Weather Underground poskytuje veřejné API pouze pro venkovní hodnoty, služba Weathercloud "
        "přijímá a uchovává i vnitřní hodnoty (teplotu v obýváku/místnosti a vlhkost).\n\n"
        "Váš systém (Cloudflare Worker) proto každých 10 minut navštíví Weathercloud a živé indoor hodnoty stáhne. "
        "Aby se Worker dostal k vašim privátním údajům, prokazuje se tzv. "
    )
    r_bold = p.add_run("session cookie")
    r_bold.bold = True
    p.add_run(
        " (přihlašovací relací z prohlížeče). Tato relace má svou platnost — obvykle funguje několik týdnů až měsíců, "
        "ale jednoho dne přirozeně vyprší."
    )
    
    add_callout(
        doc,
        "Když cookie vyprší, venkovní měření (teplota venku, vítr, déšť, tlak) funguje nerušeně dál přes Weather Underground! "
        "Pouze na kartě 'Vnitřní teplota' na webu se místo čísla objeví pomlčka '—'. "
        "Cookie tedy nemusíte obnovovat preventivně, ale až ve chvíli, kdy vnitřní teplota vypadne.",
        title="DŮLEŽITÉ: SYSTÉM ZŮSTÁVÁ BEZPEČNÝ A ŽIVÝ",
        callout_type="info"
    )
    
    # ── 2. JAK ZKONTROLOVAT STAV ZA 5 SEKUND ──────────────────────────────────
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run("2. Jak si ověřit platnost cookie na 1 klik")
    r.font.name = "Segoe UI"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(6)
    p.add_run("Vytvořili jsme pro vás speciální diagnostický odkaz. Kdykoliv máte pochybnost, otevřete v prohlížeči:")
    
    add_code_block(doc, "https://stanice-worker.petmat247.workers.dev/api/test-wc")
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Stránka okamžitě vrátí přehledný výsledek:")
    
    bp1 = doc.add_paragraph(style='List Bullet')
    bp1.paragraph_format.space_after = Pt(3)
    r1 = bp1.add_run("✅ ZELENÝ STAV (vše funguje): ")
    r1.bold = True
    r1.font.color.rgb = RGBColor(0x16, 0xA3, 0x4A)
    bp1.add_run('Uvidíte text "success": true a aktuální "tempIn": 20.2. Vše je v naprostém pořádku.')
    
    bp2 = doc.add_paragraph(style='List Bullet')
    bp2.paragraph_format.space_after = Pt(8)
    r2 = bp2.add_run("❌ ČERVENÝ STAV (cookie vypršela): ")
    r2.bold = True
    r2.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
    bp2.add_run('Uvidíte chybu se statusem 401 nebo 403 ("neplatná/vypršená session cookie"). To je přesně moment, kdy postupujte podle následujících 2 kroků.')
    
    # ── 3. KROK 1: ZÍSKÁNÍ NOVÉ COOKIE ───────────────────────────────────────
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run("3. Krok 1: Jak získat novou cookie z prohlížeče (1 minuta)")
    r.font.name = "Segoe UI"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    steps = [
        ("1. Otevřete Weathercloud a přihlaste se", "V Google Chrome nebo Safari otevřete stránku https://app.weathercloud.net a přihlaste se ke svému účtu, kde vidíte svoji stanici."),
        ("2. Otevřete Vývojářské nástroje (DevTools)", "Na Macu stiskněte klávesovou zkratku Cmd + Option + I (nebo kdekoliv na stránce klikněte pravým tlačítkem myši a zvolte 'Prozkoumat' / 'Inspect')."),
        ("3. Přepněte se na záložku 'Network' (Síť)", "V horní liště otevřeného bočního panelu klikněte na záložku Network (vedle Elements a Console)."),
        ("4. Obnovte stránku", "Stiskněte klávesu Cmd + R (obnovit stránku) na klávesnici."),
        ("5. Vyhledejte požadavek 'values'", "Do vyhledávacího políčka 'Filter' v panelu Network napište slovo: values. V seznamu pod filtrem se objeví řádek začínající 'values?code=...'."),
        ("6. Zkopírujte Cookie", "Klikněte na tento řádek. Vpravo se otevře podrobnost — záložka 'Headers'. Sjeďte o kousek dolů do sekce 'Request Headers', najděte položku 'Cookie:' a myší označte a zkopírujte (Cmd + C) celou její dlouhou hodnotu.")
    ]
    
    for s_title, s_desc in steps:
        p_step = doc.add_paragraph()
        p_step.paragraph_format.space_before = Pt(4)
        p_step.paragraph_format.space_after = Pt(3)
        p_step.paragraph_format.line_spacing = 1.15
        r_st = p_step.add_run(s_title + "\n")
        r_st.bold = True
        r_st.font.name = "Segoe UI"
        r_st.font.size = Pt(11)
        r_st.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r_sd = p_step.add_run(s_desc)
        r_sd.font.name = "Segoe UI"
        r_sd.font.size = Pt(10)
        r_sd.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    add_callout(
        doc,
        "Cookie je dlouhý řetězec obsahující písmena, čísla, podtržítka a středníky (např. remember_code=...; PHPSESSID=...; wc_session=...). "
        "Zkopírujte jej přesně tak, jak je, bez uvozovek a bez úprav.",
        title="JAK VYPADÁ COOKIE",
        callout_type="info"
    )
    
    # ── 4. KROK 2: NAHRÁNÍ NOVÉ COOKIE DO CLOUDFLARE ─────────────────────────
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run("4. Krok 2: Jak novou cookie uložit do Cloudflare")
    r.font.name = "Segoe UI"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(6)
    p.add_run("Máte na výběr ze dvou pohodlných způsobů — zvolte ten, který vám více vyhovuje:")
    
    # Varianta A
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    r = h2.add_run("Způsob A: Přes Terminál na Macu (nejrychlejší, 20 sekund)")
    r.bold = True
    r.font.name = "Segoe UI"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.add_run("1. Otevřete Terminál na Macu a zadejte:")
    add_code_block(doc, "cd ~/Desktop/Stanice\nnpx wrangler secret put WC_COOKIE --config worker/wrangler.toml")
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.add_run("2. Terminál vypíše: ")
    p.add_run("Enter a secret value:").italic = True
    p.add_run(" — stiskněte ")
    p.add_run("Cmd + V").bold = True
    p.add_run(" (vložit) a stiskněte klávesu ")
    p.add_run("Enter").bold = True
    p.add_run(".\n3. Terminál potvrdí: ")
    p.add_run("✨ Success! Uploaded secret WC_COOKIE").bold = True
    p.add_run(".\n4. Hotovo! Není potřeba spouštět žádný deploy, nová hodnota platí okamžitě.")
    
    # Varianta B
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    r = h2.add_run("Způsob B: Přímo přes web Cloudflare Dashboard (bez terminálu)")
    r.bold = True
    r.font.name = "Segoe UI"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    
    web_steps = [
        "1. Otevřete https://dash.cloudflare.com a přihlaste se.",
        "2. V levém menu klikněte na 'Workers & Pages'.",
        "3. Klikněte na název vašeho Workeru: stanice-worker.",
        "4. V horním menu záložek klikněte na 'Settings' (Nastavení).",
        "5. V levém podmenu klikněte na 'Variables and Secrets'.",
        "6. U položky WC_COOKIE klikněte na 'Edit' (nebo ikonu tužky).",
        "7. Vložte novou zkopírovanou cookie a klikněte na 'Save and Deploy' (Uložit a nasadit)."
    ]
    for ws in web_steps:
        p_ws = doc.add_paragraph()
        p_ws.paragraph_format.space_after = Pt(2)
        p_ws.paragraph_format.line_spacing = 1.15
        r_ws = p_ws.add_run(ws)
        r_ws.font.name = "Segoe UI"
        r_ws.font.size = Pt(10)
        r_ws.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    # ── 5. KROK 3: OVĚŘENÍ ───────────────────────────────────────────────────
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run("5. Krok 3: Okamžité ověření úspěchu")
    r.font.name = "Segoe UI"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(6)
    p.add_run("Nemusíte čekat 10 minut na automatický cyklus. Stačí v prohlížeči znovu kliknout na:")
    
    add_code_block(doc, "https://stanice-worker.petmat247.workers.dev/api/test-wc")
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "Tento odkaz s novou cookie ihned kontaktuje Weathercloud, stáhne aktuální teplotu a rovnou s ní aktualizuje "
        "poslední záznamy v databázi! Na webovém dashboardu "
    )
    p.add_run("https://petmat247-cloud.github.io/WeatherStation/").bold = True
    p.add_run(" okamžitě naskočí aktuální vnitřní teplota i vlhkost domova.")
    
    add_callout(
        doc,
        "Celý proces obnovy zabere méně než 2 minuty. Pokud byste cokoliv zapomněli, tento dokument můžete kdykoliv znovu otevřít ve složce docs/ nebo se podívat do POSTUP.md.",
        title="SHRNUTÍ",
        callout_type="success"
    )
    
    # Uložení dokumentu
    output_path = os.path.expanduser("/Users/petmat247/Desktop/Stanice/docs/Navod_Obnova_Weathercloud_Cookie.docx")
    doc.save(output_path)
    print(f"Dokument úspěšně vytvořen: {output_path}")

if __name__ == "__main__":
    build_tutorial_doc()
